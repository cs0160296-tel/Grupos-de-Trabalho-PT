#!/usr/bin/env python3
"""Import the unchanged workbook; detect headers, preserve cells/links, validate joins."""
from pathlib import Path
from datetime import datetime,date
from collections import Counter
import json,re,hashlib,sys,shutil
import openpyxl
ROOT=Path(__file__).resolve().parent.parent
source=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else ROOT/'data/source.xlsx'
w=openpyxl.load_workbook(source,data_only=True)
HEADERS={'Groups':['ID','Creating act'],'Deadlines':['Deadline ID','Group ID'],'Issuing portfolios':['Group ID','Normalised issuing portfolio'],'Related bodies':['Body name · Português'],'Discovery gaps':['Lead / scope record'],'Coverage':['Publication month'],'Act review':['Research classification'],'Search log':['Exact official-source outcome query'],'Methodology':['Topic','Method / limitation']}
def value(v):
 if isinstance(v,(date,datetime)):return v.strftime('%Y-%m-%d')
 return v
sheets={};header_rows={}
for name,required in HEADERS.items():
 if name not in w.sheetnames:raise ValueError('Missing sheet: '+name)
 s=w[name];found=None
 for row in s.iter_rows(min_row=1,max_row=min(30,s.max_row)):
  if all(k in [c.value for c in row] for k in required):found=row;break
 if found is None:raise ValueError('Headers not found: '+name)
 header_rows[name]=found[0].row
 headers=[c.value for c in found];records=[]
 for row in s.iter_rows(min_row=found[0].row+1):
  if all(c.value is None for c in row):continue
  raw={h:value(c.value) for h,c in zip(headers,row) if h is not None}
  links={h:c.hyperlink.target for h,c in zip(headers,row) if h and c.hyperlink and c.hyperlink.target}
  records.append({'raw':raw,'links':links})
 sheets[name]=records

def urls(*items):
 return list(dict.fromkeys(u.rstrip('.,;') for x in items if x for u in re.findall(r'https?://[^\s<>]+',str(x))))
def get(r,k):return r['raw'].get(k)
def text(r,k):return str(get(r,k) or '')
def dt(r,k):
 v=get(r,k)
 if v in (None,''):return None
 v=str(v)[:10]
 try:return date.fromisoformat(v).isoformat()
 except ValueError:raise ValueError(f'Invalid date {k}: {v}')
trans={}
for f in ['translations-a.json','translations-b.json']:
 p=ROOT/'data'/f
 if p.exists():trans.update(json.loads(p.read_text()))
# Nomes de apresentação. O nome da folha Groups é o sumário do DR: longo e por vezes
# truncado. Estes nomes servem a navegação; a designação oficial fica visível na página.
short_names=json.loads((ROOT/'data/nomes-curtos.json').read_text())
aux_path=ROOT/'data/transparency-pt.json'
aux=json.loads(aux_path.read_text()) if aux_path.exists() else {}
portfolio={}
for r in sheets['Issuing portfolios']:
 portfolio.setdefault(text(r,'Group ID'),[]).append(text(r,'Normalised issuing portfolio'))
deadlines=[]
for r in sheets['Deadlines']:
 deadlines.append({'id':text(r,'Deadline ID'),'group_id':text(r,'Group ID'),'label':text(r,'Milestone / provision'),'kind':text(r,'Type / scope'),'wording_pt':text(r,'Original Portuguese wording'),'anchor':text(r,'Trigger / anchor'),'due':dt(r,'Calculated date / benchmark'),'alternative_due':dt(r,'Alternative date'),'calculation':text(r,'Calculation / qualification'),'source_url':text(r,'Source URL'),'verification':text(r,'Verification'),'raw':r['raw']})
groups=[]
PUBLIC_REPORT_PATTERN=re.compile(r'final report is public|publication of the final report|final report publicly released|official ETF report is dated|public report,|government publishes the report',re.I)
for r in sheets['Groups']:
 gid=text(r,'ID');payment=text(r,'Participation · English')
 delivery=text(r,'Report / output evidence')
 g={'id':gid,'display_name':short_names.get(gid) or text(r,'Group name · Português'),'name':text(r,'Group name · Português'),'status':text(r,'Status · Português'),'delay_days':get(r,'Delay · days'),'government':text(r,'Government'),'act':text(r,'Creating act'),'url':r['links'].get('Creating act · click'),'published':dt(r,'DR publication date'),'signed':dt(r,'Latest signature / adoption'),'effective':dt(r,'In force / effective'),'issuer':text(r,'Issuing ministries / Secretarias de Estado · original'),'coordinator':text(r,'Coordinator'),'members':text(r,'Members · original Portuguese provisions'),'area':text(r,'Policy area'),'mission_en':text(r,'Mission · English'),'objective_pt':text(r,'Objective · original Português'),'specific_objectives':text(r,'Numbered objectives · original Português'),'deadline_original':dt(r,'Original controlling date'),'deadline_current':dt(r,'Current controlling date'),'deadline_basis':text(r,'Controlling deadline · basis'),'deadline_alternative':dt(r,'Alternative counting date'),'delay_alternative':get(r,'Alternative delay · days'),'participation':text(r,'Participation · original clause'),'participation_en':payment,'payment':('unpaid' if payment.startswith('Unpaid') else 'paid' if payment.startswith('Paid:') else 'terms' if payment.startswith('Payment') else 'unknown'),'status_explanation':text(r,'Status evidence · explanation'),'evidence_urls':urls(get(r,'Evidence URLs · includes DR'),r['links'].get('Status source · click')),'delivery_date':dt(r,'Verified delivery date'),'delivery_evidence':delivery,'has_delivery_evidence':not delivery.startswith('No '),'report_public':bool(PUBLIC_REPORT_PATTERN.search(delivery)),'report_availability_basis':delivery if PUBLIC_REPORT_PATTERN.search(delivery) else 'Não foi identificada confirmação inequívoca de um relatório final público neste registo. Pode existir trabalho entregue ou um resultado parcial; consultar a evidência.','policy_outcome':text(r,'Policy implementation · separate assessment'),'policy_evidence':urls(get(r,'Policy evidence URLs')),'extensions':text(r,'Extensions / amendments / replacements'),'has_extension':any('extension' in d['kind'].lower() for d in deadlines if d['group_id']==gid) or text(r,'Status · Português')=='Prorrogado','notes':text(r,'Notes and limitations'),'searches':text(r,'What was searched'),'last_checked':dt(r,'Last checked'),'signature_dates':text(r,'All signature dates'),'effective_basis':text(r,'Effective-date basis'),'name_basis':text(r,'Name / label basis'),'benchmark_label':text(r,'Status benchmark'),'delay_qualification':text(r,'Delay qualification'),'portfolios':list(dict.fromkeys(portfolio.get(gid,[]))),'deadlines':[d for d in deadlines if d['group_id']==gid],'raw':r['raw'],'hyperlinks':r['links']}
 fingerprints=json.loads((ROOT/'data/translation-fingerprints.json').read_text()) if (ROOT/'data/translation-fingerprints.json').exists() else {}
 fingerprint=hashlib.sha256('\n'.join(g[k] for k in ['mission_en','status_explanation','delivery_evidence','policy_outcome']).encode()).hexdigest()
 if fingerprints.get(gid)==fingerprint:g.update(trans.get(gid,{}))
 else:g['translation_pending']=True
 groups.append(g)
ids=[g['id'] for g in groups];idset=set(ids)
assert len(ids)==len(idset),'Duplicate group IDs'
assert groups,'No groups imported'
missing_names=[g['id'] for g in groups if g['id'] not in short_names]
assert not missing_names,f'Nomes de apresentação em falta em data/nomes-curtos.json: {missing_names}'
assert all(g['display_name'].startswith('Grupo de Trabalho para') for g in groups),'Nome de apresentação sem o prefixo normalizado'
assert len({g['display_name'] for g in groups})==len(groups),'Nomes de apresentação duplicados'
for g in groups:
 assert g['id'] and g['name'] and g['url'] and g['url'].startswith('https://diariodarepublica.pt/'),f'Missing creating source: {g["id"]}'
 assert g['government'] in ['XXIV','XXV']
 assert g['portfolios'],f'Missing portfolios {g["id"]}'
for d in deadlines:
 assert d['group_id'] in idset,f'Orphan deadline {d["id"]}'
 assert d['source_url'].startswith('https://'),f'Missing milestone source {d["id"]}'
assert len({d['id'] for d in deadlines})==len(deadlines),'Duplicate deadline IDs'
assert set(portfolio)<=idset,'Orphan portfolio IDs'
for name,field in [('Act review','Main-register group IDs'),('Search log','Group IDs')]:
 for r in sheets[name]:
  for gid in re.findall(r'GT-\d{4}-\d{3}',text(r,field)):assert gid in idset,f'Orphan {name}: {gid}'
cutoffs={g['last_checked'] for g in groups};assert len(cutoffs)==1,'Mixed evidence cutoffs: define a snapshot first'
cutoff=next(iter(cutoffs));assert cutoff
statuses=Counter(g['status'] for g in groups)
expected=['Concluído no prazo','Concluído com atraso','Em curso – dentro do prazo','Em atraso','Prorrogado','Extinto / substituído','Sem informação pública']
assert set(statuses)<=set(expected),'Unknown status'
overdue=[g['delay_days'] for g in groups if g['status']=='Em atraso' and g['delay_days'] is not None]
# Qualidade dos dados: anomalias do ficheiro de origem, registadas em vez de corrigidas em silêncio.
# Um prazo de entrega não pode vencer antes de o ato que o cria ser publicado. Onde isso acontece,
# a linha da folha Deadlines está a citar texto que não é uma obrigação de entrega (tipicamente uma
# alínea da composição). Corrige-se no Excel, não aqui.
published_by_group={g['id']:g['published'] for g in groups}
premature_deadlines=[{'deadline':d['id'],'group':d['group_id'],'due':d['due'],'act_published':published_by_group.get(d['group_id']),'kind':d['kind'],'wording':d['wording_pt'][:180]}
 for d in deadlines if d['due'] and published_by_group.get(d['group_id']) and d['due']<published_by_group[d['group_id']]]
# Entrega documentada mas estado não concluído: diferença entre o que a investigação encontrou e o
# que o painel publica. Ver PR de auditoria; a decisão de taxonomia é editorial.
delivery_without_completed=[g['id'] for g in groups if g['has_delivery_evidence'] and not g['status'].startswith('Concluído')]
data_quality={'deadlines_dated_before_act_publication':premature_deadlines,'delivery_evidence_without_completed_status':delivery_without_completed}
coverage=[r['raw'] for r in sheets['Coverage']]
metrics={'total':len(groups),'governments':dict(Counter(g['government'] for g in groups)),'statuses':{s:statuses[s] for s in expected},'creating_acts':len({g['url'] for g in groups}),'deadlines':len(deadlines),'average_delay':sum(overdue)/len(overdue) if overdue else None,'max_delay':max(overdue) if overdue else None,'no_information_percent':100*statuses['Sem informação pública']/len(groups),'months':len(coverage),'monthly_allocated_urls':sum(r['Distinct DR result URLs'] for r in coverage),'screened_urls':int(re.search(r'(\d+) distinct result URLs', next(r['raw']['Method / limitation'] for r in sheets['Methodology'] if r['raw']['Topic']=='Discovery search')).group(1)),'reviewed_acts':len(sheets['Act review']),'deadlines_quoted':sum(1 for d in deadlines if d['wording_pt'].strip()),'deadlines_benchmark':sum(1 for d in deadlines if not d['wording_pt'].strip()),'delivery_evidence_total':sum(1 for g in groups if g['has_delivery_evidence']),'delivery_evidence_without_completed_status':len(delivery_without_completed),'deadlines_dated_before_act_publication':len(premature_deadlines)}
# Validate the supplied historical workbook. Future dated workbooks use relational checks above.
if cutoff=='2026-10-10':
 assert metrics['total']==122 and metrics['deadlines']==441
 assert metrics['governments']=={'XXIV':74,'XXV':48}
 assert metrics['creating_acts']==116
 assert metrics['screened_urls']==3251 and metrics['reviewed_acts']==292 and metrics['months']==31
 assert metrics['statuses']==dict(zip(expected,[1,1,18,60,4,2,36]))
 assert next(g for g in groups if g['id']=='GT-2024-046')['delay_days']==102
 assert metrics['deadlines_quoted']==330 and metrics['deadlines_benchmark']==111,'Quoted/benchmark split changed'
 assert metrics['delivery_evidence_total']==20,'Delivery-evidence count changed'
 assert metrics['delivery_evidence_without_completed_status']==18,'Delivery/status gap changed'
 assert metrics['deadlines_dated_before_act_publication']==7,'Premature-deadline count changed: fix data/source.xlsx, do not relax this'
output={'cutoff':cutoff,'scope_start':'2024-04-02','groups':groups,'deadlines':deadlines,'sheets':sheets,'translations':aux,'metrics':metrics,'source':{'file':'fonte-original.xlsx','sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'header_rows':header_rows,'filename':source.name},'validation':{'unique_ids':True,'valid_sources':True,'referential_integrity':True,'translation_records':len(trans)},'data_quality':data_quality}
(ROOT/'public').mkdir(exist_ok=True)
output['dashboard_cells']=[[value(c.value) for c in row] for row in w['Dashboard'].iter_rows()]
(ROOT/'public/audit-completo.json').write_text(json.dumps(output,ensure_ascii=False,separators=(',',':')))
lean={**output,'sheets':{k:v for k,v in sheets.items() if k not in ['Groups','Deadlines','Issuing portfolios']}}
lean.pop('dashboard_cells',None)
lean['groups']=[{**g,'raw':None,'hyperlinks':None,'deadlines':[{k:v for k,v in d.items() if k!='raw'} for d in g['deadlines']]} for g in groups]
lean['deadlines']=[{k:v for k,v in d.items() if k!='raw'} for d in deadlines]
(ROOT/'public/data.json').write_text(json.dumps(lean,ensure_ascii=False,separators=(',',':')))
shutil.copyfile(source,ROOT/'public/fonte-original.xlsx')
(ROOT/'data/import-report.json').write_text(json.dumps({'cutoff':cutoff,'metrics':metrics,'validation':output['validation'],'source':output['source'],'data_quality':data_quality},ensure_ascii=False,indent=2))
if source != ROOT/'data/source.xlsx':shutil.copyfile(source,ROOT/'data/source.xlsx')
print(json.dumps({'cutoff':cutoff,**metrics,'translated':sum('mission_pt' in g for g in groups),'translations_pending':sum(g.get('translation_pending',False) for g in groups)},ensure_ascii=False))
